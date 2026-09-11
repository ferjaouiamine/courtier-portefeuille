#!/usr/bin/env python3
"""Génère la documentation Word complète du projet courtier-portefeuille."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


RACINE = Path(__file__).resolve().parents[1]
SORTIE = RACINE / "docs" / "Documentation_complete_courtier_portefeuille.docx"

VERT_FONCE = "073B34"
VERT = "0F6B5B"
VERT_CLAIR = "DDECE7"
OR = "B58200"
ROUGE = "A61B29"
GRIS = "5D6B68"
GRIS_CLAIR = "F2F5F4"
BLANC = "FFFFFF"


def couleur_cellule(cellule, couleur):
    tc_pr = cellule._tc.get_or_add_tcPr()
    fond = tc_pr.find(qn("w:shd"))
    if fond is None:
        fond = OxmlElement("w:shd")
        tc_pr.append(fond)
    fond.set(qn("w:fill"), couleur)


def marges_cellule(cellule, haut=90, debut=110, bas=90, fin=110):
    tc = cellule._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for cote, valeur in (("top", haut), ("start", debut), ("bottom", bas), ("end", fin)):
        element = tc_mar.find(qn(f"w:{cote}"))
        if element is None:
            element = OxmlElement(f"w:{cote}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(valeur))
        element.set(qn("w:type"), "dxa")


def bordures_tableau(tableau, couleur="C9D5D1", taille="4"):
    tbl_pr = tableau._tbl.tblPr
    bordures = tbl_pr.first_child_found_in("w:tblBorders")
    if bordures is None:
        bordures = OxmlElement("w:tblBorders")
        tbl_pr.append(bordures)
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        bordure = OxmlElement(f"w:{cote}")
        bordure.set(qn("w:val"), "single")
        bordure.set(qn("w:sz"), taille)
        bordure.set(qn("w:color"), couleur)
        bordures.append(bordure)


def largeur_cellule(cellule, largeur_cm):
    cellule.width = Cm(largeur_cm)
    tc_pr = cellule._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(largeur_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def garder_avec_suivant(paragraphe):
    p_pr = paragraphe._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def ajouter_numero_page(paragraphe):
    paragraphe.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraphe.add_run("Page ")
    run.font.size = Pt(8)
    champ_debut = OxmlElement("w:fldChar")
    champ_debut.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    champ_fin = OxmlElement("w:fldChar")
    champ_fin.set(qn("w:fldCharType"), "end")
    run._r.extend([champ_debut, instruction, champ_fin])


def ajouter_toc(document):
    paragraphe = document.add_paragraph()
    run = paragraphe.add_run()
    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separation = OxmlElement("w:fldChar")
    separation.set(qn("w:fldCharType"), "separate")
    texte = OxmlElement("w:t")
    texte.text = "Ouvrez le document dans Word puis actualisez la table des matières."
    separation.append(texte)
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run._r.extend([debut, instruction, separation, fin])


def ajouter_hyperlien(paragraphe, texte, url):
    relation_id = paragraphe.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlien = OxmlElement("w:hyperlink")
    hyperlien.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    propriete = OxmlElement("w:rPr")
    couleur = OxmlElement("w:color")
    couleur.set(qn("w:val"), VERT)
    soulignement = OxmlElement("w:u")
    soulignement.set(qn("w:val"), "single")
    propriete.extend([couleur, soulignement])
    run.append(propriete)
    contenu = OxmlElement("w:t")
    contenu.text = texte
    run.append(contenu)
    hyperlien.append(run)
    paragraphe._p.append(hyperlien)


def ajouter_tableau(document, entetes, lignes, largeurs=None):
    tableau = document.add_table(rows=1, cols=len(entetes))
    tableau.alignment = WD_TABLE_ALIGNMENT.CENTER
    tableau.autofit = True
    bordures_tableau(tableau)
    for index, entete in enumerate(entetes):
        cellule = tableau.rows[0].cells[index]
        couleur_cellule(cellule, VERT_FONCE)
        marges_cellule(cellule)
        cellule.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cellule.paragraphs[0].add_run(str(entete))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
        if largeurs:
            largeur_cellule(cellule, largeurs[index])
    tr_pr = tableau.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    for numero, ligne in enumerate(lignes):
        cellules = tableau.add_row().cells
        for index, valeur in enumerate(ligne):
            cellule = cellules[index]
            marges_cellule(cellule)
            cellule.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if numero % 2:
                couleur_cellule(cellule, GRIS_CLAIR)
            paragraphe = cellule.paragraphs[0]
            paragraphe.paragraph_format.space_after = Pt(0)
            run = paragraphe.add_run(str(valeur))
            run.font.size = Pt(8.5)
            if largeurs:
                largeur_cellule(cellule, largeurs[index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return tableau


def ajouter_encadre(document, titre, texte, couleur=VERT_CLAIR, accent=VERT):
    tableau = document.add_table(rows=1, cols=1)
    tableau.alignment = WD_TABLE_ALIGNMENT.CENTER
    cellule = tableau.cell(0, 0)
    couleur_cellule(cellule, couleur)
    marges_cellule(cellule, 140, 160, 140, 160)
    bordures_tableau(tableau, accent, "8")
    paragraphe = cellule.paragraphs[0]
    run = paragraphe.add_run(titre)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(accent)
    paragraphe.add_run("\n" + texte)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def ajouter_code(document, lignes):
    tableau = document.add_table(rows=1, cols=1)
    cellule = tableau.cell(0, 0)
    couleur_cellule(cellule, "17211F")
    marges_cellule(cellule, 110, 140, 110, 140)
    paragraphe = cellule.paragraphs[0]
    paragraphe.paragraph_format.space_after = Pt(0)
    run = paragraphe.add_run(lignes)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(235, 244, 241)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def ajouter_liste(document, elements, numerotee=False):
    style = "List Number" if numerotee else "List Bullet"
    for element in elements:
        paragraphe = document.add_paragraph(style=style)
        paragraphe.add_run(element)


def titre_section(document, texte, niveau=1):
    paragraphe = document.add_heading(texte, level=niveau)
    garder_avec_suivant(paragraphe)
    return paragraphe


def ajouter_separateur(document):
    paragraphe = document.add_paragraph()
    p_pr = paragraphe._p.get_or_add_pPr()
    bordures = OxmlElement("w:pBdr")
    bas = OxmlElement("w:bottom")
    bas.set(qn("w:val"), "single")
    bas.set(qn("w:sz"), "10")
    bas.set(qn("w:color"), OR)
    bordures.append(bas)
    p_pr.append(bordures)


def configurer_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("243330")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for nom, taille, couleur in (
        ("Title", 34, VERT_FONCE),
        ("Subtitle", 16, GRIS),
        ("Heading 1", 20, VERT_FONCE),
        ("Heading 2", 14, VERT),
        ("Heading 3", 11, OR),
    ):
        style = styles[nom]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(taille)
        style.font.color.rgb = RGBColor.from_string(couleur)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if nom != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    if "Légende technique" not in [style.name for style in styles]:
        style = styles.add_style("Légende technique", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style.font.size = Pt(8)
        style.font.italic = True
        style.font.color.rgb = RGBColor.from_string(GRIS)


def configurer_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    entete = section.header.paragraphs[0]
    entete.text = "FINASURE  |  COURTIER PORTEFEUILLE"
    entete.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    entete.runs[0].font.size = Pt(8)
    entete.runs[0].font.bold = True
    entete.runs[0].font.color.rgb = RGBColor.from_string(VERT)

    pied = section.footer.paragraphs[0]
    pied.add_run("Documentation interne  •  7 septembre 2026                         ")
    pied.runs[0].font.size = Pt(8)
    pied.runs[0].font.color.rgb = RGBColor.from_string(GRIS)
    ajouter_numero_page(pied)

    document.core_properties.title = "Documentation complète - Courtier Portefeuille"
    document.core_properties.subject = "Documentation fonctionnelle, technique, exploitation et pérennisation"
    document.core_properties.author = "Finasure Solutions"
    document.core_properties.keywords = "assurance, portefeuille, contrats, échéances, Neon, Vercel"

    # Demande à Word de recalculer la table des matières et les numéros de page.
    settings = document.settings._element
    mise_a_jour = OxmlElement("w:updateFields")
    mise_a_jour.set(qn("w:val"), "true")
    settings.append(mise_a_jour)


def generer():
    document = Document()
    configurer_styles(document)
    configurer_document(document)

    # Couverture
    document.add_paragraph().paragraph_format.space_after = Pt(55)
    marque = document.add_paragraph()
    marque.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = marque.add_run("FINASURE SOLUTIONS")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(VERT)

    titre = document.add_paragraph(style="Title")
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titre.add_run("Courtier Portefeuille")
    sous_titre = document.add_paragraph(style="Subtitle")
    sous_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_titre.add_run("Documentation fonctionnelle, technique, exploitation et plan de pérennisation")
    ajouter_separateur(document)

    infos = document.add_paragraph()
    infos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    infos.add_run("Version documentaire 1.0\n").bold = True
    infos.add_run("État du code : commit 95b6a5c\n")
    infos.add_run("Date de référence : 7 septembre 2026\n")
    ajouter_hyperlien(infos, "https://courtier-portefeuille.vercel.app", "https://courtier-portefeuille.vercel.app")

    document.add_paragraph().paragraph_format.space_after = Pt(45)
    ajouter_encadre(
        document,
        "CONFIDENTIALITÉ",
        "Document interne. Il ne contient volontairement aucun mot de passe, jeton JWT, jeton Blob ni URL PostgreSQL réelle. Les secrets doivent rester exclusivement dans les variables d’environnement sécurisées.",
        "F7E9E9",
        ROUGE,
    )
    document.add_page_break()

    titre_section(document, "Fiche du document")
    ajouter_tableau(
        document,
        ["Élément", "Valeur"],
        [
            ("Produit", "Application web de gestion de portefeuille d’assurance"),
            ("Propriétaire fonctionnel", "Finasure Solutions"),
            ("Dépôt", "ferjaouiamine/courtier-portefeuille"),
            ("Environnement public", "Vercel, domaine courtier-portefeuille.vercel.app"),
            ("Base de données", "PostgreSQL hébergé sur Neon"),
            ("Documents", "Vercel Blob privé en production, disque local en développement"),
            ("Public visé", "Administrateurs, agents, lecteurs et futurs profils comptables"),
            ("Révision recommandée", "Après chaque changement métier majeur et au minimum chaque trimestre"),
        ],
        [4.2, 12.5],
    )

    titre_section(document, "Historique documentaire", 2)
    ajouter_tableau(
        document,
        ["Version", "Date", "Objet", "Statut"],
        [("1.0", "07/09/2026", "Première documentation consolidée à partir du code", "Référence initiale")],
        [2, 3, 8.5, 3.2],
    )

    titre_section(document, "Table des matières")
    ajouter_toc(document)
    ajouter_encadre(
        document,
        "MISE À JOUR DANS WORD",
        "Cliquez dans la table des matières, puis choisissez « Mettre à jour toute la table » pour recalculer les numéros de page.",
    )
    document.add_page_break()

    # 1
    titre_section(document, "1. Résumé exécutif")
    document.add_paragraph(
        "Courtier Portefeuille centralise les clients, les contrats d’assurance, les échéances, les encaissements, les relances, les pièces jointes et les référentiels d’un cabinet de courtage. L’application remplace un suivi dispersé dans Excel par une base PostgreSQL structurée, une interface responsive et une traçabilité nominative des opérations."
    )
    document.add_paragraph(
        "La solution actuelle convient à une petite équipe et possède déjà plusieurs fondations utiles pour évoluer : isolation par organisation, pagination, index PostgreSQL, contrôle de rôles, journal d’audit et stockage objet privé. Pour devenir un service durable utilisé par plusieurs milliers de comptes, elle doit encore renforcer la gestion des identités, les migrations, les sauvegardes cloud, l’observabilité, les tests et les traitements asynchrones."
    )
    ajouter_tableau(
        document,
        ["Axe", "État actuel", "Objectif long terme"],
        [
            ("Métier", "Parcours principaux opérationnels", "Règles validées, paramétrables et documentées"),
            ("Sécurité", "JWT, cookie httpOnly, bcrypt, rôles et RLS", "MFA, révocation de sessions, contrôle fin et audit immuable"),
            ("Scalabilité", "Pagination et index présents", "Pool maîtrisé, cache, files de travaux et tests de charge"),
            ("Exploitation", "Vercel + Neon + Blob", "Supervision, alertes, sauvegarde vérifiée et procédure d’incident"),
            ("Qualité", "Tests unitaires et quelques intégrations", "CI complète, tests E2E et déploiement contrôlé"),
        ],
        [3, 6, 7.5],
    )
    ajouter_encadre(
        document,
        "PRIORITÉ IMMÉDIATE",
        "Faire tourner les secrets déjà partagés dans des conversations ou captures, vérifier le stockage Blob privé, définir une stratégie de sauvegarde Neon et contrôler les données historiques après le passage au calcul automatique de la date de fin.",
        "FFF4D6",
        OR,
    )

    # 2
    titre_section(document, "2. Périmètre fonctionnel")
    titre_section(document, "2.1 Fonctionnalités disponibles", 2)
    ajouter_liste(document, [
        "Connexion sécurisée par e-mail et mot de passe, avec session en cookie httpOnly.",
        "Tableau de bord : contrats, primes, échéances en retard, échéances sous 30 jours, répartition par compagnie et branche, encaissements sur 12 mois.",
        "Frise et liste des échéances des 120 prochains jours ; un clic ouvre le contrat concerné.",
        "Gestion des clients personnes physiques ou morales, avec date de naissance pour les personnes physiques.",
        "Création, consultation, modification, recherche, tri, archivage, restauration et suppression définitive des contrats et clients selon les règles de dépendance.",
        "Gestion du souscripteur, du payeur et d’une société de leasing saisie en texte libre.",
        "Référentiel des compagnies et produits, classés par branche d’assurance.",
        "Échéances filtrables par fenêtre, urgence, type, compagnie et recherche.",
        "Encaissement total ou partiel et recalcul automatique du statut de l’échéance.",
        "Relances manuelles par appel, SMS, WhatsApp, e-mail ou type automatique.",
        "Pièces jointes PDF, JPG et PNG rattachées aux contrats.",
        "Export CSV du portefeuille, compatible avec Excel en français.",
        "Journal des activités indiquant l’administrateur, l’action, la date et les champs modifiés.",
        "Corbeille avec restauration et suppression définitive réservées aux administrateurs.",
        "Import initial depuis un fichier Excel avec rapport des corrections et anomalies.",
    ])

    titre_section(document, "2.2 Hors périmètre ou incomplet aujourd’hui", 2)
    ajouter_liste(document, [
        "Aucune interface autonome de gestion des utilisateurs, des rôles ou des mots de passe.",
        "Aucun parcours « mot de passe oublié », MFA ou fermeture immédiate de toutes les sessions.",
        "Le rôle comptable et le rôle super-administrateur sont proposés dans ce document mais ne sont pas encore implémentés.",
        "Le script npm de relances automatiques référence un fichier absent ; l’automatisation des relances n’est donc pas opérationnelle.",
        "La date de fin reste stockée techniquement mais elle est calculée automatiquement à partir de la date d’effet et de la durée.",
        "Les sauvegardes fournies ciblent Docker local ; la production Neon/Blob exige une procédure cloud distincte.",
        "Les pièces jointes envoyées par le serveur Vercel sont limitées à 4 Mo.",
        "Les champs historiques de commission restent présents dans le schéma même s’ils ne font plus partie du parcours métier principal ; une décision de migration est nécessaire.",
    ])

    # 3
    titre_section(document, "3. Utilisateurs et droits")
    ajouter_tableau(
        document,
        ["Rôle", "Droits actuels", "Restrictions"],
        [
            ("admin", "Lecture et écriture ; journal ; corbeille ; restauration ; suppression définitive", "Ne doit pas disposer du mot de passe de la base de migration"),
            ("agent", "Lecture ; ajout et modification ; archivage ; paiements ; relances ; pièces jointes", "Pas de journal global, restauration ni suppression définitive"),
            ("lecture", "Consultation des données de son organisation", "Aucune écriture"),
        ],
        [2.5, 8, 6],
    )

    titre_section(document, "3.1 Proposition cible", 2)
    ajouter_tableau(
        document,
        ["Rôle cible", "Responsabilité recommandée"],
        [
            ("super_admin", "Administration technique des organisations et comptes ; aucun usage métier quotidien ; MFA obligatoire."),
            ("admin_metier", "Gestion complète du portefeuille de son organisation, sauf configuration technique globale."),
            ("comptable", "Lecture clients/contrats, saisie et correction contrôlée des encaissements, exports et rapprochements ; pas de suppression de référentiel."),
            ("agent", "Gestion commerciale et suivi des échéances sans suppression définitive."),
            ("lecture", "Consultation et export selon autorisation."),
        ],
        [3.5, 13],
    )
    ajouter_encadre(
        document,
        "PRINCIPE DU MOINDRE PRIVILÈGE",
        "Un compte nominatif par personne, jamais de compte partagé. Le super-administrateur est réservé aux opérations exceptionnelles ; les activités quotidiennes utilisent un rôle moins puissant.",
    )

    # 4
    titre_section(document, "4. Règles métier")
    titre_section(document, "4.1 Client, souscripteur, assuré, payeur et leasing", 2)
    document.add_paragraph(
        "Le client principal est la personne rattachée au contrat. Par défaut, le souscripteur est le client principal et le payeur est le souscripteur. Lorsque le souscripteur diffère de l’assuré, il faut sélectionner la fiche correspondante. La société de leasing est un champ texte libre : par exemple, un véhicule peut être assuré au nom de Walid, avec CIL comme société de leasing."
    )
    ajouter_tableau(
        document,
        ["Situation", "Saisie recommandée"],
        [
            ("Particulier payant son contrat", "Client = souscripteur = payeur ; société de leasing vide."),
            ("Véhicule financé par leasing", "Client/assuré = personne couverte ; souscripteur selon le contrat ; société de leasing = raison sociale saisie librement."),
            ("Entreprise payant pour une personne", "Client/assuré et souscripteur selon le document ; payeur = fiche entreprise."),
        ],
        [5.5, 11],
    )

    titre_section(document, "4.2 Durée, dates et fréquence de paiement", 2)
    document.add_paragraph(
        "La durée métier possède deux valeurs : Durée ferme et Renouvelable par tacite reconduction (RTR). Une durée ferme impose une prime unique et ne génère aucune échéance. Un contrat RTR autorise uniquement une fréquence trimestrielle, semestrielle ou annuelle."
    )
    ajouter_tableau(
        document,
        ["Fractionnement", "Calcul actuel", "Exemple pour effet au 10/01/2027"],
        [
            ("Trimestriel", "Date d’effet + 3 mois", "10/04/2027"),
            ("Semestriel", "Date d’effet + 6 mois", "10/07/2027"),
            ("Annuel", "Date d’effet + 12 mois", "10/01/2028"),
            ("Prime unique", "Aucune échéance de terme générée", "Aucune échéance"),
        ],
        [4, 6, 6.5],
    )
    document.add_paragraph(
        "Pour un contrat RTR, la date d’échéance métier est calculée automatiquement à partir de la date d’effet et de la fréquence. Un encaissement total clôture le terme courant et crée le terme suivant ; un encaissement partiel conserve le terme courant. Les dates restent ancrées sur la date d’effet pour éviter une dérive en fin de mois. Les échéanciers personnalisés ne sont jamais prolongés automatiquement."
    )
    ajouter_encadre(
        document,
        "RÈGLE AUTOMATIQUE",
        "L’utilisateur choisit Durée ferme ou RTR. Pour une durée ferme, la date de fin est obligatoire et doit être postérieure à la date d’effet. Pour un contrat RTR, la date de fin reste une donnée technique non modifiable.",
    )

    titre_section(document, "4.3 Niveaux et statuts d’échéance", 2)
    ajouter_tableau(
        document,
        ["État", "Définition"],
        [
            ("À venir", "Aucun paiement et date d’échéance non dépassée."),
            ("Partielle", "Un montant positif a été encaissé, mais le total reste inférieur à la prime."),
            ("Payée", "Somme des paiements actifs supérieure ou égale au montant de l’échéance."),
            ("Impayée", "Aucun paiement suffisant et date d’échéance strictement antérieure à aujourd’hui."),
            ("En retard", "Niveau d’affichage pour une échéance non payée dont la date est passée."),
            ("Urgente", "Échéance non payée dans 15 jours ou moins."),
            ("À préparer", "Échéance non payée dans 45 jours ou moins, hors urgence/retard."),
            ("À jour", "Échéance plus éloignée."),
        ],
        [4, 12.5],
    )

    titre_section(document, "4.4 Tableau de bord", 2)
    ajouter_liste(document, [
        "Total contrats et contrats en cours : uniquement les contrats et clients non archivés.",
        "Primes émises : somme des primes des contrats en cours.",
        "En retard et sous 30 jours : échéances de terme non payées.",
        "Frise et liste 120 jours : échéances de terme comprises entre aujourd’hui et aujourd’hui + 120 jours, dates incluses.",
        "Encaissements 12 mois : paiements actifs depuis le premier jour du mois situé onze mois avant le mois courant jusqu’à aujourd’hui ; les contrats, clients et échéances archivés sont exclus.",
        "Un clic sur une échéance proche ouvre directement la fiche du contrat.",
    ])

    titre_section(document, "4.5 Archivage et suppression", 2)
    document.add_paragraph(
        "L’archivage est une suppression logique : la ligne reste en base avec supprime_le et supprime_par mais disparaît des vues actives et des indicateurs. La restauration est réservée à l’administrateur. La suppression définitive efface les dépendances d’un contrat (relances, paiements, pièces jointes et échéances), puis le contrat. Un client ne peut pas être supprimé définitivement tant qu’un contrat le référence."
    )
    ajouter_encadre(
        document,
        "ATTENTION",
        "La suppression définitive est irréversible hors sauvegarde. La procédure cible doit imposer une confirmation renforcée, une durée minimale en corbeille et, pour les données sensibles, une validation par un second administrateur.",
        "F7E9E9",
        ROUGE,
    )

    # 5
    titre_section(document, "5. Architecture technique")
    ajouter_tableau(
        document,
        ["Couche", "Technologie", "Responsabilité"],
        [
            ("Navigateur", "HTML, CSS, JavaScript natif", "Interface responsive, formulaires, tableaux, navigation SPA et appels API."),
            ("Application", "Node.js 20+, Express 4", "Authentification, autorisations, validations, règles métier et API REST."),
            ("Données", "PostgreSQL / Neon", "Persistance, contraintes, vues, fonctions, triggers, RLS et audit."),
            ("Documents", "Vercel Blob privé / S3 / local", "Stockage des PDF et images ; les métadonnées restent en PostgreSQL."),
            ("Hébergement", "Vercel", "Build et exécution de l’application, variables sécurisées et déploiement depuis GitHub."),
            ("Développement local", "Docker Compose", "PostgreSQL local et application reproductible."),
        ],
        [3, 4.5, 9],
    )

    titre_section(document, "5.1 Flux principal", 2)
    ajouter_code(document, "Navigateur authentifié\n        │ HTTPS + cookie httpOnly\n        ▼\nVercel / Express API\n        ├──► Neon PostgreSQL pooler\n        │      données + RLS + audit\n        └──► Vercel Blob privé\n               PDF / JPG / PNG")
    document.add_paragraph(
        "Chaque requête authentifiée récupère organisationId dans le JWT. AsyncLocalStorage transporte ce contexte jusqu’à la couche SQL, qui positionne app.organisation_id dans la transaction. Les politiques RLS PostgreSQL filtrent ensuite les lignes de l’organisation courante."
    )

    titre_section(document, "5.2 Modèle de données", 2)
    ajouter_tableau(
        document,
        ["Table", "Contenu", "Relations majeures"],
        [
            ("organisations", "Cabinets isolés", "Parent logique de toutes les données métier."),
            ("utilisateurs", "Comptes, rôles, hachages", "Rattachés à une organisation."),
            ("clients", "Personnes physiques/morales", "Client, souscripteur, leasing historique ou payeur d’un contrat."),
            ("compagnies", "Compagnies d’assurance", "Référencées par les contrats."),
            ("produits", "Produits et branches", "Référencés par les contrats."),
            ("contrats", "Police, dates, fréquence, prime, acteurs", "Parent des échéances et pièces jointes."),
            ("echeances", "Terme ou renouvellement, date, montant, statut", "Parent des paiements et relances."),
            ("paiements", "Encaissements", "Rattachés à une échéance."),
            ("relances", "Actions de suivi", "Rattachées à une échéance."),
            ("pieces_jointes_contrats", "Métadonnées de fichiers", "Pointe vers l’objet Blob/S3 par nom_stockage."),
            ("journal_audit", "Avant/après, auteur, action, date", "Traçabilité transversale par organisation."),
        ],
        [4, 6, 6.5],
    )

    titre_section(document, "5.3 Index et pagination déjà présents", 2)
    ajouter_liste(document, [
        "Recherche trigramme sur le nom du client.",
        "Index par organisation et nom de client.",
        "Index contrats par organisation/statut/date d’effet et par numéro de contrat.",
        "Index échéances par organisation/date/statut.",
        "Index paiements par organisation/date et par échéance.",
        "Index pièces jointes par contrat/date d’ajout.",
        "Pagination serveur sur clients, contrats, échéances et journal, limitée à 100 lignes par page.",
    ])

    # 6
    titre_section(document, "6. Installation locale")
    titre_section(document, "6.1 Prérequis", 2)
    ajouter_liste(document, [
        "Git.",
        "Node.js 20 ou supérieur et npm.",
        "Docker Desktop pour PostgreSQL local recommandé.",
        "Python et openpyxl uniquement pour régénérer un import Excel.",
    ])

    titre_section(document, "6.2 Installation avec Docker", 2)
    ajouter_code(document, "git clone https://github.com/ferjaouiamine/courtier-portefeuille.git\ncd courtier-portefeuille\nnpm ci\nCopy-Item .env.example .env\ndocker compose up -d db")
    document.add_paragraph(
        "Dans .env, DATABASE_URL doit utiliser le rôle applicatif courtier_app. MIGRATION_DATABASE_URL doit utiliser le propriétaire local PostgreSQL pour appliquer le schéma. Utiliser des mots de passe distincts."
    )
    ajouter_code(document, "npm run migrate\nnpm run creer-admin -- \"Administrateur Local\" admin.local@exemple.tn \"MOT_DE_PASSE_SOLIDE\"\nnpm run dev")
    document.add_paragraph("Ouvrir http://localhost:3000 puis vérifier http://localhost:3000/api/sante.")

    titre_section(document, "6.3 Commandes utiles Windows", 2)
    ajouter_code(document, "# Trouver le PID qui écoute sur le port 3000\nnetstat -ano | findstr LISTENING | findstr :3000\n\n# Arrêter le processus identifié\ntaskkill /PID NUMERO_PID /F\n\n# Relancer l'application\nnpm run dev")
    document.add_paragraph(
        "Une ligne TIME_WAIT avec PID 0 n’est pas un serveur à arrêter. Il faut chercher une ligne LISTENING."
    )

    # 7
    titre_section(document, "7. Configuration des environnements")
    ajouter_tableau(
        document,
        ["Variable", "Obligatoire", "Usage et règle"],
        [
            ("DATABASE_URL", "Oui", "URL PostgreSQL du rôle applicatif, sans SUPERUSER ni BYPASSRLS ; utiliser le pooler Neon en production."),
            ("MIGRATION_DATABASE_URL", "Migration seulement", "Compte propriétaire du schéma ; ne pas l’exposer au serveur web si possible."),
            ("JWT_SECRET", "Oui", "Secret aléatoire long, différent par environnement ; jamais égal au mot de passe PostgreSQL."),
            ("JWT_DUREE", "Non", "Durée des sessions, 12h par défaut."),
            ("NODE_ENV", "Oui en prod", "Valeur production pour activer les cookies Secure."),
            ("PORT", "Local", "Port local ; Vercel le gère automatiquement."),
            ("STOCKAGE_DRIVER", "Oui en prod", "local, vercel-blob ou s3."),
            ("BLOB_READ_WRITE_TOKEN", "Blob classique", "Jeton créé par la connexion Vercel Blob ; secret absolu."),
            ("S3_* / AWS_*", "Si S3", "Bucket, région, endpoint éventuel et identifiants à privilèges minimaux."),
            ("CONNEXION_*", "Non", "Seuil et durée du blocage temporaire de connexion."),
            ("SMS_*", "Si SMS", "Configuration du fournisseur ; automatisation à finaliser."),
        ],
        [4.4, 2.6, 9.5],
    )
    ajouter_encadre(
        document,
        "GESTION DES SECRETS",
        "Ne jamais mettre .env dans Git, Word, e-mail ou messagerie. Tout secret visible dans une capture ou une conversation doit être révoqué et remplacé. Générer un JWT_SECRET avec : node -e \"console.log(require('crypto').randomBytes(48).toString('base64url'))\".",
        "F7E9E9",
        ROUGE,
    )

    # 8
    titre_section(document, "8. Déploiement Vercel, Neon et Blob")
    titre_section(document, "8.1 Préparer Neon", 2)
    ajouter_liste(document, [
        "Créer ou sélectionner le projet PostgreSQL dans une région européenne proche des utilisateurs et de Vercel.",
        "Conserver un rôle propriétaire réservé aux migrations.",
        "Créer un rôle courtier_app dédié à l’API, sans privilège BYPASSRLS.",
        "Appliquer le schéma avec MIGRATION_DATABASE_URL depuis une machine d’administration contrôlée.",
        "Tester la connexion applicative et vérifier /api/sante.",
    ], numerotee=True)
    ajouter_code(document, "$env:MIGRATION_DATABASE_URL=\"postgresql://ROLE_MIGRATION:SECRET@HOTE/neondb?sslmode=require\"\n$env:DATABASE_URL=\"postgresql://courtier_app:SECRET@HOTE-POOLER/neondb?sslmode=require\"\nnpm run migrate")
    document.add_paragraph(
        "Ces commandes se lancent dans PowerShell, jamais dans l’éditeur SQL Neon. Dans l’éditeur SQL, écrire uniquement du SQL."
    )

    titre_section(document, "8.2 Configurer Vercel", 2)
    ajouter_liste(document, [
        "Importer le dépôt GitHub et conserver la racine du projet.",
        "Définir DATABASE_URL, JWT_SECRET, JWT_DUREE, NODE_ENV=production et STOCKAGE_DRIVER=vercel-blob.",
        "Ne pas définir PORT manuellement.",
        "Limiter MIGRATION_DATABASE_URL aux opérations de migration ; idéalement ne pas la conserver dans les variables d’exécution Vercel.",
        "Déployer puis contrôler /api/sante et la page de connexion.",
    ], numerotee=True)

    titre_section(document, "8.3 Configurer les contrats PDF", 2)
    ajouter_liste(document, [
        "Dans le projet Vercel, ouvrir Storage puis Create Database et Blob.",
        "Choisir une région européenne et l’accès Private.",
        "Utiliser le préfixe de variables BLOB, sans tiret.",
        "Activer l’ajout du jeton lecture/écriture afin d’obtenir BLOB_READ_WRITE_TOKEN.",
        "Vérifier STOCKAGE_DRIVER=vercel-blob pour Production et Preview.",
        "Redéployer, puis tester l’ajout, le téléchargement et la suppression d’un PDF inférieur ou égal à 4 Mo.",
    ], numerotee=True)
    ajouter_encadre(
        document,
        "CONFIDENTIALITÉ DES CONTRATS",
        "Le store doit rester Private. Le téléchargement traverse l’API authentifiée : ne jamais remplacer cette logique par une URL Blob publique pour des contrats clients.",
    )

    titre_section(document, "8.4 Déploiement standard", 2)
    ajouter_code(document, "git status\nnpm ci\nnpm test\ngit add <fichiers>\ngit commit -m \"Description claire\"\ngit push origin main")
    document.add_paragraph(
        "Vercel redéploie automatiquement après le push. Contrôler le déploiement de production, pas seulement une URL de Preview protégée."
    )

    # 9
    titre_section(document, "9. Utilisation quotidienne")
    titre_section(document, "9.1 Nouveau client", 2)
    ajouter_liste(document, [
        "Rechercher d’abord par nom, CIN/matricule ou téléphone pour éviter un doublon.",
        "Créer la fiche et choisir personne physique ou morale.",
        "Pour une personne physique, renseigner la date de naissance ; une date future est refusée.",
        "Renseigner le CIN/matricule, le téléphone et le code Finasure lorsque disponibles.",
        "Ouvrir la fiche pour contrôler les contrats liés et l’historique.",
    ], numerotee=True)

    titre_section(document, "9.2 Nouveau contrat", 2)
    ajouter_liste(document, [
        "Vérifier le client et créer au préalable les éventuels souscripteur et payeur.",
        "Renseigner numéro, compagnie, produit, prime et date d’effet, puis choisir Durée ferme ou RTR.",
        "Pour une durée ferme, renseigner une date de fin postérieure à la date d’effet ; aucune échéance n’est générée.",
        "Pour un contrat RTR, choisir une fréquence annuelle, semestrielle ou trimestrielle et contrôler la date d’échéance calculée.",
        "Saisir la société de leasing dans le champ texte lorsqu’elle existe.",
        "Enregistrer, ouvrir la fiche et joindre le contrat signé en PDF.",
        "Vérifier que l’échéance apparaît dans l’agenda et, si elle tombe dans les 120 jours, sur le tableau de bord.",
    ], numerotee=True)

    titre_section(document, "9.3 Encaissement", 2)
    ajouter_liste(document, [
        "Ouvrir Échéances et filtrer selon la fenêtre ou le niveau.",
        "Cliquer sur Encaisser, saisir montant, date, mode et référence.",
        "Vérifier le statut recalculé : partielle ou payée.",
        "Après un encaissement total RTR, contrôler la date du prochain terme affichée par l’application.",
        "Ne pas créer de faux paiement pour masquer une échéance en retard.",
    ], numerotee=True)

    titre_section(document, "9.4 Archivage", 2)
    ajouter_liste(document, [
        "Archiver d’abord les contrats concernés avant d’archiver un client qui possède des contrats en cours.",
        "Contrôler le tableau de bord : les contrats, échéances et encaissements archivés doivent disparaître des indicateurs.",
        "Restaurer depuis la corbeille en cas d’erreur.",
        "N’utiliser la suppression définitive qu’après expiration du délai interne de rétention et contrôle de la sauvegarde.",
    ], numerotee=True)

    # 10
    titre_section(document, "10. Sécurité actuelle et renforcements")
    titre_section(document, "10.1 Protections présentes", 2)
    ajouter_liste(document, [
        "Mots de passe hachés avec bcrypt, facteur de coût 12.",
        "JWT signé, conservé dans un cookie httpOnly, SameSite=Strict et Secure en production.",
        "Limitation générale de l’API et limitation spécifique de la connexion.",
        "Rôles admin, agent et lecture vérifiés côté serveur.",
        "Requêtes SQL paramétrées.",
        "Isolation par organisation via Row-Level Security forcée dans PostgreSQL.",
        "Contrôle des références inter-organisations par triggers.",
        "Journal d’audit avec utilisateur et états avant/après.",
        "Stockage privé des pièces jointes.",
    ])

    titre_section(document, "10.2 Risques à corriger", 2)
    ajouter_tableau(
        document,
        ["Priorité", "Risque", "Mesure recommandée"],
        [
            ("P0", "Secrets exposés ou réutilisés", "Rotation immédiate de la base, du JWT, du Blob et des mots de passe administrateurs."),
            ("P0", "Pas de restauration cloud testée", "Sauvegarde logique externe et exercice de restauration documenté."),
            ("P1", "Blocage de connexion en mémoire", "Stockage distribué Redis/PostgreSQL pour fonctionner sur plusieurs instances serverless."),
            ("P1", "Sessions JWT non révocables", "Sessions serveur ou version de jeton, rotation et déconnexion globale."),
            ("P1", "Pas de MFA ni reset sécurisé", "MFA obligatoire pour admins et flux de récupération à jeton court."),
            ("P1", "CSP désactivée", "Activer une Content-Security-Policy testée et retirer les exceptions inutiles."),
            ("P1", "Fichiers validés seulement par MIME déclaré", "Vérifier signature réelle, analyser antivirus et mettre en quarantaine."),
            ("P1", "Audit annoncé append-only mais droits SQL génériques", "Révoquer UPDATE/DELETE à l’application sur journal_audit et protéger par trigger/role dédié."),
            ("P2", "Pas de CSRF explicite", "Ajouter un jeton CSRF pour les écritures, en complément de SameSite."),
            ("P2", "Validation API dispersée", "Adopter des schémas centralisés avec Zod/Joi et limites de chaînes."),
        ],
        [1.5, 5.5, 9.5],
    )

    # 11
    titre_section(document, "11. Sauvegarde, restauration et continuité")
    titre_section(document, "11.1 Objectifs à faire valider", 2)
    ajouter_tableau(
        document,
        ["Indicateur", "Cible initiale proposée", "Signification"],
        [
            ("RPO", "24 heures, puis 1 heure", "Perte maximale de données acceptable après incident."),
            ("RTO", "4 heures", "Temps maximal pour restaurer le service."),
            ("Rétention", "30 jours quotidiens + 12 mensuelles", "À adapter aux obligations contractuelles et légales."),
        ],
        [3, 4, 9.5],
    )

    titre_section(document, "11.2 Développement Docker", 2)
    ajouter_code(document, "npm run backup\n\n# Restauration destructive après validation du dossier\npowershell -ExecutionPolicy Bypass -File scripts/restore.ps1 `\n  -BackupDirectory \"C:\\sauvegardes\\AAAAmmjj-HHMMSS\" -Force")
    document.add_paragraph(
        "Le script calcule une empreinte SHA-256 du dump et archive les pièces jointes locales. Il est conçu pour Docker local, pas pour Neon."
    )

    titre_section(document, "11.3 Production Neon et Blob", 2)
    ajouter_liste(document, [
        "Activer et comprendre la restauration temporelle proposée par le plan Neon utilisé.",
        "Produire un pg_dump chiffré vers un compte de sauvegarde distinct de Neon.",
        "Exporter régulièrement l’inventaire des objets Blob et conserver une seconde copie chiffrée des contrats critiques.",
        "Tester chaque mois une restauration dans un projet Neon isolé, jamais directement sur la production.",
        "Après restauration, contrôler /api/sante, les comptes, les totaux clients/contrats, les échéances, les encaissements et plusieurs fichiers.",
        "Tracer date, opérateur, durée, résultat et anomalies de chaque exercice.",
    ], numerotee=True)

    titre_section(document, "11.4 Procédure d’incident", 2)
    ajouter_liste(document, [
        "Qualifier l’incident : disponibilité, corruption, fuite, erreur humaine ou fournisseur.",
        "Geler les écritures si leur poursuite aggrave la situation.",
        "Préserver logs, horodatages et identifiants de déploiement.",
        "Choisir le dernier point de restauration conforme au RPO.",
        "Restaurer dans un environnement isolé et effectuer les contrôles métier.",
        "Basculer seulement après validation par un responsable métier et un responsable technique.",
        "Documenter la cause racine et les actions empêchant la récidive.",
    ], numerotee=True)

    # 12
    titre_section(document, "12. Tests et processus de livraison")
    document.add_paragraph(
        "La suite actuelle utilise node:test et supertest. Elle couvre les autorisations, la propagation de l’organisation, l’isolation RLS sur base locale, l’audit et certains effets d’archivage. Les tests destructifs sont volontairement ignorés sur une base distante."
    )
    ajouter_code(document, "npm ci\nnpm test\nnode --check src/server.js\ngit diff --check")
    titre_section(document, "12.1 Pipeline CI cible", 2)
    ajouter_liste(document, [
        "Installation reproductible avec npm ci.",
        "Analyse statique, formatage et audit des dépendances.",
        "Tests unitaires sans base.",
        "Tests d’intégration sur PostgreSQL éphémère avec migration complète.",
        "Tests E2E Playwright : connexion, client, contrat, échéance, paiement, archivage et PDF.",
        "Test explicite de fuite entre deux organisations.",
        "Build et Preview Vercel.",
        "Migration de production exécutée comme étape contrôlée et journalisée.",
        "Smoke tests sur /api/sante et les parcours essentiels.",
        "Possibilité de revenir au déploiement précédent sans annuler une migration destructive non compatible.",
    ], numerotee=True)

    titre_section(document, "12.2 Règle de migration", 2)
    document.add_paragraph(
        "Le schéma actuel est un fichier SQL idempotent sans historique de versions. À long terme, adopter des migrations numérotées et immuables. Utiliser la stratégie expand/contract : ajouter d’abord les nouveaux champs compatibles, déployer le code, migrer les données, puis supprimer les anciens champs dans une version ultérieure."
    )

    # 13
    titre_section(document, "13. Supervision et exploitation")
    ajouter_tableau(
        document,
        ["Signal", "Seuil ou contrôle", "Action"],
        [
            ("Disponibilité", "/api/sante toutes les minutes", "Alerter après plusieurs échecs consécutifs."),
            ("Erreurs HTTP", "Taux 5xx et 429", "Examiner logs Vercel avec correlation_id."),
            ("Latence", "p50, p95, p99 des API", "Analyser requêtes lentes et saturation du pool."),
            ("PostgreSQL", "Connexions, CPU, stockage, requêtes lentes", "Ajuster pool/index et archiver si nécessaire."),
            ("Blob", "Volume, erreurs put/get/delete", "Vérifier token, quota et cohérence DB/objets."),
            ("Métier", "Échéances en retard et relances non traitées", "Alerter le responsable opérationnel."),
            ("Sauvegarde", "Dernier succès et dernier test de restauration", "Incident immédiat si fenêtre RPO dépassée."),
        ],
        [3.5, 5, 8],
    )
    ajouter_liste(document, [
        "Ajouter un identifiant de corrélation à chaque requête et le retourner dans les erreurs.",
        "Centraliser les erreurs avec Sentry ou un outil équivalent, sans données personnelles ni secrets.",
        "Configurer une surveillance externe de disponibilité.",
        "Créer un tableau de bord technique séparé du tableau de bord métier.",
        "Définir un responsable d’astreinte et une procédure d’escalade.",
    ])

    # 14
    titre_section(document, "14. Plan de montée en charge")
    document.add_paragraph(
        "Quelques milliers d’utilisateurs sont réalistes si la montée en charge est pilotée. Le facteur déterminant n’est pas seulement le nombre de comptes, mais le nombre de requêtes simultanées, le volume des contrats, la fréquence des tableaux de bord et la taille des documents."
    )
    ajouter_tableau(
        document,
        ["Étape", "Déclencheur", "Actions"],
        [
            ("Niveau 1", "Équipe interne, trafic faible", "Conserver l’architecture ; sécuriser secrets, sauvegardes, migrations et monitoring."),
            ("Niveau 2", "Centaines d’utilisateurs actifs", "Pool Neon borné, rate limit distribué, cache des référentiels, index issus de mesures, jobs asynchrones."),
            ("Niveau 3", "Milliers d’utilisateurs", "Tests de charge, files de messages, tableaux de bord pré-agrégés, quotas par organisation, observabilité complète."),
            ("Niveau 4", "Forte croissance ou exigences fortes", "Séparer workers/API, stratégie multi-région, réplication, plan de reprise et revue d’architecture formelle."),
        ],
        [2.5, 4.5, 9.5],
    )

    titre_section(document, "14.1 Base de données", 2)
    ajouter_liste(document, [
        "Utiliser systématiquement le pooler Neon pour DATABASE_URL et limiter le nombre de connexions pg par instance.",
        "Ajouter PG_POOL_MAX, délais de connexion et délais de requête configurables.",
        "Mesurer avec pg_stat_statements avant d’ajouter des index.",
        "Paginer toutes les listes et éviter les exports synchrones sans limite.",
        "Pré-agréger le tableau de bord si son calcul devient coûteux.",
        "Mettre en place des quotas et tests d’isolation pour chaque nouvelle table.",
    ])

    titre_section(document, "14.2 Traitements et fichiers", 2)
    ajouter_liste(document, [
        "Déplacer relances, imports, exports volumineux et rapports vers des workers idempotents.",
        "Utiliser une file avec retries, temporisation et file d’échec.",
        "Pour des documents supérieurs à 4 Mo, utiliser un upload direct Blob avec jeton court signé par l’API.",
        "Analyser les fichiers antivirus avant disponibilité et imposer une politique de rétention.",
        "Contrôler périodiquement les objets orphelins et les métadonnées sans objet.",
    ])

    # 15
    titre_section(document, "15. Feuille de route priorisée")
    ajouter_tableau(
        document,
        ["Horizon", "Priorités", "Critère de sortie"],
        [
            ("0–30 jours", "Rotation secrets ; Blob privé ; sauvegarde Neon ; contrôle des dates historiques ; réparer relances ; CI minimale ; monitoring disponibilité", "Aucun secret exposé, restauration testée et parcours critiques surveillés."),
            ("1–3 mois", "Gestion utilisateurs ; rôle comptable ; super-admin séparé ; MFA ; reset mot de passe ; migrations versionnées ; E2E", "Droits validés automatiquement et comptes administrables sans SQL."),
            ("3–6 mois", "Rate limit distribué ; sessions révocables ; CSP/CSRF ; antivirus ; logs structurés ; jobs asynchrones", "Application sûre sur plusieurs instances et traitements fiables."),
            ("6–12 mois", "Tests de charge ; pré-agrégations ; PRA complet ; quotas ; audit externe ; politique de données", "Capacité et reprise démontrées par des exercices mesurés."),
        ],
        [2.6, 8.2, 5.7],
    )

    titre_section(document, "15.1 Backlog technique explicite", 2)
    ajouter_tableau(
        document,
        ["ID", "Action", "Priorité", "Effort indicatif"],
        [
            ("SEC-01", "Faire tourner tous les secrets et mots de passe exposés", "P0", "0,5 jour"),
            ("OPS-01", "Sauvegarde Neon + restauration de preuve", "P0", "1–2 jours"),
            ("MET-01", "Contrôler et, si nécessaire, recalculer les dates techniques des anciens contrats", "P0", "1–2 jours"),
            ("REL-01", "Créer et tester scripts/relances.js", "P1", "2–4 jours"),
            ("IAM-01", "Écran utilisateurs + rôles comptable/super_admin", "P1", "5–10 jours"),
            ("IAM-02", "MFA, reset et révocation de sessions", "P1", "5–10 jours"),
            ("DB-01", "Migrations numérotées et pipeline", "P1", "3–5 jours"),
            ("OBS-01", "Logs structurés, erreurs et alertes", "P1", "2–4 jours"),
            ("FIL-01", "Validation binaire, antivirus et upload direct", "P1", "4–8 jours"),
            ("QA-01", "Tests PostgreSQL éphémères et Playwright", "P1", "5–8 jours"),
            ("PERF-01", "Pool, tests de charge et optimisation mesurée", "P2", "4–8 jours"),
            ("DATA-01", "Décider puis migrer les champs de commission historiques", "P2", "1–3 jours"),
        ],
        [1.5, 8.5, 2.2, 4.3],
    )

    # 16
    titre_section(document, "16. Diagnostic des erreurs fréquentes")
    ajouter_tableau(
        document,
        ["Symptôme", "Cause probable", "Correction"],
        [
            ("Base de données injoignable", "DATABASE_URL incorrecte, rôle/mot de passe, SSL ou Neon suspendu", "Tester la chaîne hors application, vérifier le pooler et redéployer après correction."),
            ("401 sur /api/auth/moi", "Absence, expiration ou invalidation du cookie JWT", "Se reconnecter ; vérifier JWT_SECRET stable entre déploiements."),
            ("500 sur connexion", "Erreur SQL, schéma incomplet ou audit sans organisation", "Lire les logs Vercel et réappliquer le schéma avec le rôle de migration."),
            ("permission denied for schema public", "Migration lancée avec courtier_app", "Utiliser MIGRATION_DATABASE_URL du propriétaire pour npm run migrate."),
            ("syntax error near npm", "Commande npm saisie dans l’éditeur SQL", "Exécuter npm dans PowerShell ; réserver l’éditeur Neon au SQL."),
            ("insecure password", "Mot de passe de rôle refusé par Neon", "Générer un secret long et aléatoire avec plusieurs classes de caractères."),
            ("organisation_id null dans journal_audit", "Contexte organisation absent ou schéma incomplet", "Vérifier JWT organisationId, transaction et migration du schéma."),
            ("Ajout PDF impossible", "Blob absent, driver local sur Vercel, token manquant ou fichier > 4 Mo", "Store Blob Private, préfixe BLOB, jeton RW, STOCKAGE_DRIVER=vercel-blob, redéploiement."),
            ("Port 3000 occupé", "Ancien processus en écoute", "Identifier LISTENING avec netstat puis taskkill /PID ... /F."),
            ("src refspec main does not match any", "Aucun commit local ou branche différente", "Créer le premier commit, renommer la branche main puis pousser."),
        ],
        [4.2, 5.8, 7.2],
    )

    # 17
    titre_section(document, "17. Contrôles réguliers")
    ajouter_tableau(
        document,
        ["Fréquence", "Contrôles"],
        [
            ("Chaque jour", "Disponibilité, erreurs 5xx, échecs de relance, échéances en retard, succès des sauvegardes."),
            ("Chaque semaine", "Dépendances, stockage Blob, utilisateurs actifs, anomalies d’audit et requêtes lentes."),
            ("Chaque mois", "Restauration de test, revue des accès admin, quotas, coûts et cohérence DB/Blob."),
            ("Chaque trimestre", "Rotation planifiée, revue de sécurité, capacité, RPO/RTO et mise à jour du présent document."),
            ("Chaque année", "Test complet du PRA, audit externe, revue des durées de conservation et des fournisseurs."),
        ],
        [3.5, 13],
    )

    titre_section(document, "17.1 Checklist avant mise en production", 2)
    ajouter_liste(document, [
        "Tests automatisés réussis sur une base éphémère.",
        "Migration répétée avec succès sur une copie de production.",
        "Sauvegarde récente et restauration vérifiée.",
        "Variables Production correctes, sans secret dans Git.",
        "Compte courtier_app sans privilèges de migration ni BYPASSRLS.",
        "Blob Private connecté et test des trois opérations fichier.",
        "Smoke tests connexion, client, contrat, échéance, paiement et archivage.",
        "Plan de retour arrière défini.",
        "Responsable métier informé de la fenêtre de changement.",
    ])

    # 18
    titre_section(document, "18. API principale")
    ajouter_tableau(
        document,
        ["Méthode", "Route", "Usage"],
        [
            ("GET", "/api/sante", "Santé de l’application et de PostgreSQL."),
            ("POST", "/api/auth/connexion", "Connexion et création du cookie."),
            ("POST", "/api/auth/deconnexion", "Déconnexion."),
            ("GET", "/api/auth/moi", "Session courante."),
            ("GET/POST", "/api/clients", "Liste paginée et création."),
            ("GET/PUT/DELETE", "/api/clients/:id", "Fiche, modification et archivage."),
            ("POST", "/api/clients/:id/restaurer", "Restauration admin."),
            ("GET/POST", "/api/contrats", "Portefeuille paginé et création."),
            ("GET/PUT/DELETE", "/api/contrats/:id", "Fiche, modification et archivage."),
            ("POST", "/api/contrats/:id/renouveler", "Renouvellement."),
            ("POST/GET/DELETE", "/api/contrats/:id/pieces-jointes/...", "Ajouter, télécharger et supprimer un fichier."),
            ("GET", "/api/echeances", "Agenda paginé et filtré."),
            ("POST", "/api/echeances/:id/paiements", "Encaissement."),
            ("POST", "/api/echeances/:id/relances", "Relance."),
            ("GET", "/api/tableau-de-bord", "Indicateurs et échéances proches."),
            ("GET", "/api/journal-audit", "Journal global admin."),
            ("GET/DELETE", "/api/corbeille/...", "Liste et suppression définitive admin."),
            ("GET", "/api/export/portefeuille.csv", "Export CSV."),
        ],
        [2, 7, 7.5],
    )

    # 19
    titre_section(document, "19. Gouvernance et décisions à formaliser")
    ajouter_liste(document, [
        "Qui est propriétaire fonctionnel des règles d’échéance et des exceptions ?",
        "Qui valide une suppression définitive et après quel délai ?",
        "Quelle durée légale et contractuelle de conservation s’applique aux contrats, paiements et journaux ?",
        "Quelles opérations exactes sont autorisées au futur rôle comptable ?",
        "Quel RPO/RTO est accepté et quel budget fournisseur en découle ?",
        "Les commissions historiques doivent-elles être conservées, exportées puis supprimées du schéma ?",
        "Quelle taille maximale de pièce jointe est réellement nécessaire ?",
        "Quels canaux de relance sont autorisés et comment recueillir le consentement ?",
        "Qui reçoit les alertes et qui décide du retour arrière d’un déploiement ?",
    ])
    ajouter_encadre(
        document,
        "RÈGLE DE GOUVERNANCE",
        "Chaque décision doit avoir un propriétaire, une date, une justification et un critère d’acceptation. Les règles métier ne doivent pas être modifiées uniquement dans le code sans validation et mise à jour documentaire.",
    )

    # 20
    titre_section(document, "20. Glossaire")
    ajouter_tableau(
        document,
        ["Terme", "Définition"],
        [
            ("Échéance", "Date calculée à laquelle un règlement ou une action métier devient attendu."),
            ("Fractionnement", "Fréquence contractuelle utilisée ici pour calculer la date d’échéance : 3, 6 ou 12 mois ; prime unique sans échéance."),
            ("RLS", "Row-Level Security : filtrage des lignes imposé par PostgreSQL selon l’organisation."),
            ("JWT", "Jeton signé décrivant la session ; ce n’est ni un mot de passe ni une URL de base."),
            ("Blob", "Objet binaire stocké séparément de PostgreSQL, par exemple un PDF."),
            ("RPO", "Quantité maximale de données qu’une restauration peut faire perdre."),
            ("RTO", "Durée maximale acceptable pour rétablir le service."),
            ("PRA", "Plan de reprise d’activité après incident majeur."),
            ("Migration", "Modification versionnée et reproductible du schéma de base."),
            ("Idempotent", "Traitement pouvant être rejoué sans produire de doublons ou incohérences."),
        ],
        [4, 12.5],
    )

    document.add_page_break()
    fin = document.add_paragraph(style="Title")
    fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fin.add_run("Fin du document")
    conclusion = document.add_paragraph()
    conclusion.alignment = WD_ALIGN_PARAGRAPH.CENTER
    conclusion.add_run(
        "Ce document décrit l’état observé du code au 7 septembre 2026. Toute évolution métier, de sécurité ou d’infrastructure doit entraîner une révision de cette documentation."
    )

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    document.save(SORTIE)
    print(SORTIE)


if __name__ == "__main__":
    generer()
